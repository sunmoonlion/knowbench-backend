"""读用户的数据：分页、抽文字、按格子读、核引文（SDD 0042；随读数层从 knowlab 搬来，原 0027 §3.5、
0030 §3.1、0036）。knowlab 那一道核与跨应用契约的测试留在原处，这里只有纯逻辑与抽取。
"""

from __future__ import annotations

import asyncio
import datetime as dt
import hashlib
import io
import json
import threading

import pytest

from app.domain.workbench.data_read import (
    PAGE_BYTES,
    Extracted,
    Sheet,
    TableUnreadable,
    find_quote,
    normalize_quote,
    paginate,
    pdf_pages,
    table_page,
)
from app.infrastructure.documents import table_extract, text_extract
from app.infrastructure.documents.table_extract import (
    CachedTableExtractor,
    _sheet,
    _xls_value,
    extract_tables,
    number_of,
)
from app.infrastructure.documents.text_extract import (
    NO_TEXT,
    UNREADABLE,
    CachedTextExtractor,
    extract_text,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _bytes(data: bytes):  # type: ignore[no-untyped-def]
    async def load() -> bytes:
        return data

    return load


def test_paginate_splits_at_the_page_size() -> None:
    text = "x" * (PAGE_BYTES * 2 + 5)  # 一行就超了：按字切开
    assert paginate(text, 1) == ("x" * PAGE_BYTES, 3)
    assert paginate(text, 3) == ("x" * 5, 3)
    assert paginate("", 1) == ("", 1)


def test_a_chinese_annual_report_is_paged_by_bytes_and_keeps_its_pages_whole() -> None:
    """中文年报按字节切（待办账 137）。

    每页都在模型看得到的上限里，在「## Page N」前面切，接起来就是原文。
    """
    printed = [
        f"## Page {n}\n"
        + "".join(f"营业收入{n}-{i}，170,899,152,276.34 元\n" for i in range(400))
        for n in range(1, 13)
    ]
    text = "".join(printed)
    first, pages = paginate(text, 1)
    parts = [paginate(text, p)[0] for p in range(1, pages + 1)]
    assert "".join(parts) == text
    assert all(len(part.encode("utf-8")) <= PAGE_BYTES for part in parts)
    assert all(part.startswith("## Page ") for part in parts)  # 没把年报的一页拆开
    assert pages > 1 and paginate(text, pages + 1) == ("", pages)


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "利润表"
    sheet.append(["年份", "营收", None])
    sheet.append([2025, 100.0, None])
    book.create_sheet("说明").append(["单位：万元"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _pdf(text: str) -> bytes:
    """最小的一页 PDF，内容是一行 ASCII 字。"""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"<< /Size {len(objects) + 1} /Root 1 0 R >>"
    out += f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_excel_becomes_csv_per_sheet() -> None:
    got = extract_text(_xlsx(), XLSX)
    assert (
        got.text
        == "## Sheet: 利润表\n年份,营收\n2025,100\n\n## Sheet: 说明\n单位：万元\n\n"
    )
    assert not got.truncated and got.note is None


def test_pdf_text_comes_out_page_by_page() -> None:
    got = extract_text(_pdf("Revenue 100"), "application/pdf")
    assert got.text.startswith("## Page 1\nRevenue 100")
    assert got.note is None


def test_a_pdf_without_text_says_so() -> None:
    got = extract_text(_pdf(""), "application/pdf")
    assert got.note == NO_TEXT


def test_broken_files_are_unreadable_not_errors() -> None:
    assert extract_text(b"%PDF-1.4 broken", "application/pdf").note == UNREADABLE
    assert extract_text(b"PK\x03\x04junk", XLSX).note == UNREADABLE
    assert extract_text(b"x", "image/png").note == UNREADABLE


def test_a_workbook_that_unpacks_too_large_is_not_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(text_extract, "MAX_XLSX_UNZIPPED", 10)
    assert extract_text(_xlsx(), XLSX).note == UNREADABLE


def test_text_is_cut_at_the_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(text_extract, "MAX_TEXT_CHARS", 10)
    got = extract_text(("一" * 50).encode(), "text/plain; charset=utf-8")
    assert got.text == "一" * 10 and got.truncated and got.note


async def test_the_same_content_is_extracted_once(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[str] = []

    def fake(data: bytes, content_type: str) -> Extracted:
        calls.append(content_type)
        return Extracted(text=data.decode())

    monkeypatch.setattr(text_extract, "extract_text", fake)
    extractor = CachedTextExtractor()
    sha = hashlib.sha256(b"abc").hexdigest()
    got = await asyncio.gather(
        *[
            extractor.extract(
                load=_bytes(b"abc"), content_type="text/plain", sha256=sha
            )
            for _ in range(3)
        ]
    )
    assert [g.text for g in got] == ["abc"] * 3 and calls == ["text/plain"]


async def test_a_reader_that_gives_up_still_leaves_the_text_cached(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """抽得慢、读的那头先超时走了：抽取照样做完进缓存，下一次直接拿到（K19 预演）。"""
    calls: list[str] = []
    gate = threading.Event()

    def slow(data: bytes, content_type: str) -> Extracted:
        calls.append(content_type)
        gate.wait(5)
        return Extracted(text=data.decode())

    monkeypatch.setattr(text_extract, "extract_text", slow)
    extractor = CachedTextExtractor()
    sha = hashlib.sha256(b"abc").hexdigest()
    with pytest.raises(TimeoutError):
        await asyncio.wait_for(
            extractor.extract(
                load=_bytes(b"abc"), content_type="application/pdf", sha256=sha
            ),
            timeout=0.05,
        )
    gate.set()
    for _ in range(100):  # 等后台那次抽完
        if not extractor._pending:
            break
        await asyncio.sleep(0.01)
    got = await extractor.extract(
        load=_bytes(b"abc"), content_type="application/pdf", sha256=sha
    )
    assert got.text == "abc" and calls == ["application/pdf"]


# ---------- 按格子读 ----------


def _statements() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "利润表"
    sheet.append(["科目", dt.date(2024, 3, 31), dt.datetime(2024, 6, 30), None])
    sheet.append(["营业收入", 46485000000.12, "1,234.50", None])
    sheet.append(["营业成本", "(3,000)", "−12", None])
    sheet.append([None, None, None, None])
    sheet.append(["其中：利息", 0.1, "-", None])
    sheet.append(["其中：利息", "12.5%", "", None])
    book.create_sheet("说明").append(["单位：元"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_excel_becomes_labelled_cells_per_sheet() -> None:
    tables = extract_tables(_statements(), XLSX)
    sheet = tables.sheet("利润表")
    assert sheet is not None and tables.sheet(None) == sheet
    assert sheet.columns == ("2024-03-31", "2024-06-30")
    assert [r.label for r in sheet.rows] == [
        "营业收入",
        "营业成本",
        "其中：利息",
        "其中：利息#2",
    ]
    income = sheet.row("营业收入")
    assert income is not None
    assert income.cells["2024-03-31"].number == "46485000000.12"
    assert (income.cells["2024-06-30"].raw, income.cells["2024-06-30"].number) == (
        "1,234.50",
        "1234.5",
    )
    cost = sheet.row("营业成本")
    assert cost is not None
    assert [c.number for c in cost.cells.values()] == ["-3000", "-12"]
    interest = sheet.row("其中：利息")
    assert interest is not None and interest.cells["2024-03-31"].number == "0.1"
    assert interest.cells["2024-06-30"].number is None
    second = sheet.row("其中：利息#2")
    assert (
        second is not None and second.cells["2024-03-31"].number is None
    )  # 百分号不猜
    assert [s.name for s in tables.sheets] == ["利润表", "说明"]


def test_repeated_column_labels_are_numbered_not_overwritten() -> None:
    """合并与母公司并排、同一报告期出现两次：后一列不能盖掉前一列（审读第 5 条）。"""
    tables = extract_tables(
        "科目,2024-12-31,2024-12-31,,2023-12-31\n营收,100,60,x,90\n".encode(),
        "text/csv",
    )
    sheet = tables.sheet(None)
    assert sheet is not None
    assert sheet.columns == ("2024-12-31", "2024-12-31#2", "2023-12-31")
    row = sheet.row("营收")
    assert row is not None
    assert {k: c.number for k, c in row.cells.items()} == {
        "2024-12-31": "100",
        "2024-12-31#2": "60",
        "2023-12-31": "90",
    }


def test_xls_date_cells_become_dates_not_serial_numbers() -> None:
    import xlrd

    class XlsCell:
        def __init__(self, ctype: int, value: object) -> None:
            self.ctype, self.value = ctype, value

    assert _xls_value(XlsCell(xlrd.XL_CELL_DATE, 45657.0), 0) == dt.datetime(
        2024, 12, 31
    )
    assert _xls_value(XlsCell(xlrd.XL_CELL_NUMBER, 45657.0), 0) == 45657.0
    assert _xls_value(XlsCell(xlrd.XL_CELL_DATE, 1e20), 0) == 1e20  # 坏日期照原值
    header = [_xls_value(XlsCell(xlrd.XL_CELL_TEXT, "科目"), 0)] + [
        _xls_value(XlsCell(xlrd.XL_CELL_DATE, 45657.0), 0)
    ]
    sheet = _sheet("利润表", [header, ["营收", 100.0]])
    assert sheet.columns == ("2024-12-31",)


def test_csv_is_one_table_named_data() -> None:
    tables = extract_tables(
        "\ufeff科目,2024\n营收,100\n".encode(), "text/csv; charset=utf-8"
    )
    assert [(s.name, s.columns) for s in tables.sheets] == [("data", ("2024",))]


@pytest.mark.parametrize(
    ("value", "number"),
    [
        (100, "100"),
        (1.5, "1.5"),
        (float("nan"), None),
        (float("inf"), None),
        (True, None),
        ("1e400", None),
        ("0.00", "0"),
        ("-0", "0"),
        ("（12）", None),
        ("abc", None),
        ("--", None),
        (dt.date(2024, 1, 1), None),
    ],
)
def test_numbers_are_read_conservatively(value: object, number: str | None) -> None:
    assert number_of(value) == number


def test_non_tables_and_broken_files_are_unreadable() -> None:
    for data, kind in [
        (b"x", "text/plain"),
        (b"PK\x03\x04junk", XLSX),
        (b"junk", "application/vnd.ms-excel"),
        (b"%PDF-1.4", "application/pdf"),
    ]:
        with pytest.raises(TableUnreadable):
            extract_tables(data, kind)


def test_tables_have_size_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(table_extract, "MAX_XLSX_UNZIPPED", 10)
    with pytest.raises(TableUnreadable):
        extract_tables(_statements(), XLSX)
    monkeypatch.setattr(table_extract, "MAX_ROWS", 1)
    with pytest.raises(TableUnreadable):
        extract_tables(b"k,v\na,1\nb,2\n", "text/csv")


def test_a_table_page_holds_about_four_thousand_cells() -> None:
    from app.domain.workbench.data_read import Row

    rows = tuple(Row(label=str(i), cells={}) for i in range(450))
    sheet = Sheet(name="s", columns=tuple(str(i) for i in range(10)), rows=rows)
    first, pages = table_page(sheet, 1)
    assert (len(first), pages) == (400, 2)
    assert len(table_page(sheet, 2)[0]) == 50 and table_page(sheet, 3)[0] == ()
    assert table_page(Sheet(name="e", columns=(), rows=()), 1) == ((), 1)


def test_a_wide_table_page_stays_within_the_bytes_the_model_can_see() -> None:
    """格子不到 4000，但字节超了：按字节再分页（待办账 137）。"""
    from app.domain.workbench.data_read import Cell, Row

    columns = tuple(f"2025-{m:02d}-30" for m in range(1, 13))
    rows = tuple(
        Row(
            label=f"归属于母公司所有者权益（或股东权益）合计{i}",
            cells={c: Cell(raw="170,899,152,276.34") for c in columns},
        )
        for i in range(300)
    )
    sheet = Sheet(name="资产负债表", columns=columns, rows=rows)
    _, pages = table_page(sheet, 1)
    assert pages > 1
    shown = [table_page(sheet, p)[0] for p in range(1, pages + 1)]
    assert sum(len(part) for part in shown) == 300
    for part in shown:
        body = {
            "columns": list(columns),
            "rows": [
                {"label": r.label, "values": {c: x.raw for c, x in r.cells.items()}}
                for r in part
            ],
        }
        assert len(json.dumps(body, ensure_ascii=False).encode("utf-8")) <= PAGE_BYTES


async def test_the_same_table_is_parsed_once(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    real = table_extract.extract_tables

    def counting(data: bytes, content_type: str):  # type: ignore[no-untyped-def]
        calls.append(content_type)
        return real(data, content_type)

    monkeypatch.setattr(table_extract, "extract_tables", counting)
    extractor = CachedTableExtractor()
    data = b"k,v\na,1\n"
    sha = hashlib.sha256(data).hexdigest()
    loads: list[int] = []

    async def load() -> bytes:
        loads.append(1)
        return data

    for _ in range(3):
        await extractor.tables(load=load, content_type="text/csv", sha256=sha)
    assert calls == ["text/csv"] and loads == [1]


# ---------- 核引文（SDD 0036） ----------

REPORT = (
    "## Page 1\n封面\n\n"
    "## Page 2\n第三节 管理层讨论与分析\n"
    "报告期内，公司营业收入同比增长\n15.7%，主要\n\n"
    "## Page 3\n得益于系列酒销量增加；毛利率略有下降。\n\n"
)


def test_pdf_pages_follow_the_extraction_markers() -> None:
    pages = pdf_pages(REPORT)
    assert sorted(pages) == [1, 2, 3]
    assert "管理层讨论与分析" in pages[2] and "封面" not in pages[2]
    assert pdf_pages("没有页的文字") == {}


def test_quotes_match_after_normalizing_spaces_and_widths() -> None:
    pages = pdf_pages(REPORT)
    assert (
        normalize_quote(" 营业收入 同比\n增长　１５．７％ ") == "营业收入同比增长15.7%"
    )
    assert find_quote(pages, 2, "公司营业收入同比增长 15.7%") is None
    assert find_quote(pages, 2, "公司营业收入同比增长１５．７％") is None  # 全角
    assert (
        find_quote(pages, 2, "公司营业收入同比增长 15.8%") == "not_found"
    )  # 改了一个字
    assert find_quote(pages, 3, "公司营业收入同比增长 15.7%") == "not_found"  # 页码不对


def test_a_quote_may_run_onto_the_next_page_but_must_start_on_its_own() -> None:
    pages = pdf_pages(REPORT)
    assert find_quote(pages, 2, "15.7%，主要得益于系列酒销量增加") is None
    # 整段都在第 3 页：写成第 2 页不算
    assert find_quote(pages, 2, "得益于系列酒销量增加；毛利率略有下降") == "not_found"
    assert find_quote(pages, 3, "得益于系列酒销量增加；毛利率略有下降") is None


def test_quote_misses_say_why() -> None:
    pages = pdf_pages(REPORT)
    assert find_quote(pages, 2, "营业收入") == "too_short"
    assert find_quote(pages, 2, "字" * 1001) == "too_long"
    assert find_quote({}, 1, "公司营业收入同比增长") == "no_pages"
    assert find_quote(pages, 9, "公司营业收入同比增长") == "page_out_of_range"


def test_spaces_between_letters_and_digits_still_count() -> None:
    """去掉所有空白会让「now here」等于「nowhere」、「1 234」等于「12 34」。

    审读 10-19。
    """
    pages = {1: "This is now here in the report. Annual sales: 1 234 million yuan."}
    assert find_quote(pages, 1, "This is now here in the report.") is None
    assert find_quote(pages, 1, "This is nowhere in the report.") == "not_found"
    assert find_quote(pages, 1, "Annual sales: 12 34 million") == "not_found"
    assert (
        find_quote(pages, 1, "Annual  sales:\n1 234 million") is None
    )  # 断行、多空格照认


def test_numbers_alone_are_too_short_to_prove_a_passage() -> None:
    pages = {1: "截至2024年12月31日，营业收入为1,234,567.89万元"}
    assert find_quote(pages, 1, "2024年12月31日") == "too_short"
    assert find_quote(pages, 1, "1,234,567.89万元") == "too_short"
    assert (
        find_quote(pages, 1, "截至2024年12月31日，营业收入为1,234,567.89万元") is None
    )


def test_a_quote_past_a_truncated_last_page_is_out_of_range() -> None:
    pages = {1: "第一页的叙述性文字一段", 2: "被截断的最后一页只抽到了前面"}
    quote = "最后一页后面没抽到的那一段原话"
    assert find_quote(pages, 2, quote) == "not_found"
    assert find_quote(pages, 2, quote, truncated=True) == "page_out_of_range"
    assert find_quote(pages, 1, quote, truncated=True) == "not_found"
